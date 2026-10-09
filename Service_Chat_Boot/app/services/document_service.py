from __future__ import annotations

import logging
from datetime import datetime, timezone
from pathlib import Path

from fastapi import UploadFile

from app.config.settings import get_settings
from app.schemas.platform import DocumentInfo
from app.utils.exceptions import AppError

logger = logging.getLogger(__name__)

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".doc", ".xlsx", ".xls"}


def get_kb_path() -> Path:
    settings = get_settings()
    kb = Path(settings.knowledge_base_dir)
    if not kb.is_absolute():
        kb = Path(__file__).resolve().parent.parent.parent / kb
    kb.mkdir(parents=True, exist_ok=True)
    return kb


def _safe_filename(name: str) -> str:
    return Path(name).name.replace("..", "").strip()


def list_documents() -> list[DocumentInfo]:
    kb = get_kb_path()
    files: list[DocumentInfo] = []
    for path in sorted(kb.iterdir()):
        if not path.is_file() or path.name.startswith("."):
            continue
        stat = path.stat()
        files.append(
            DocumentInfo(
                name=path.name,
                size_bytes=stat.st_size,
                modified_at=datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat(),
            )
        )
    return files


async def save_upload(file: UploadFile) -> str:
    if not file.filename:
        raise AppError("Nom de fichier manquant", status_code=400)
    filename = _safe_filename(file.filename)
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise AppError(
            f"Extension non supportée: {ext}. Autorisé: {', '.join(sorted(ALLOWED_EXTENSIONS))}",
            status_code=400,
        )
    kb = get_kb_path()
    dest = kb / filename
    data = await file.read()
    dest.write_bytes(data)
    logger.info("Uploaded document: %s (%d bytes)", filename, len(data))
    return filename


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            parts.append(text)
    return "\n".join(parts)


def _extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"--- {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def _extract_xls(path: Path) -> str:
    raise AppError("Les fichiers .xls ne sont pas supportés. Convertissez en .xlsx.", status_code=400)


def _extract_doc(path: Path) -> str:
    raise AppError("Les fichiers .doc ne sont pas supportés. Convertissez en .docx.", status_code=400)


def extract_text_from_file(path: Path) -> str:
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            return _extract_pdf(path)
        if ext == ".docx":
            return _extract_docx(path)
        if ext == ".doc":
            return _extract_doc(path)
        if ext == ".xlsx":
            return _extract_xlsx(path)
        if ext == ".xls":
            return _extract_xls(path)
    except AppError:
        raise
    except Exception as exc:
        logger.warning("Failed to extract %s: %s", path.name, exc)
        return ""
    return ""


def extract_text_from_kb() -> str:
    kb = get_kb_path()
    parts: list[str] = []
    for path in sorted(kb.iterdir()):
        if not path.is_file() or path.suffix.lower() not in ALLOWED_EXTENSIONS:
            continue
        if path.suffix.lower() in {".doc", ".xls"}:
            continue
        text = extract_text_from_file(path)
        if text.strip():
            parts.append(f"=== {path.name} ===\n{text.strip()}")
    return "\n\n".join(parts)
